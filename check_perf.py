import openpyxl
import glob
import os

# 查找文件
desktop = r'C:\Users\Administrator\Desktop'
files = glob.glob(os.path.join(desktop, '运营绩效*.xlsx'))
print("找到的文件:", files)

if not files:
    print("未找到文件")
    exit()

wb = openpyxl.load_workbook(files[0])
print(f"\n工作表：{wb.sheetnames}")

ws = wb.active
print(f"\n当前工作表：{ws.title}")

# 输出所有单元格内容
print("\n=== 表格原始内容 ===")
for row_idx, row in enumerate(ws.iter_rows(), 1):
    cells = []
    for cell in row:
        if cell.value is not None:
            cells.append(f"{cell.coordinate}={repr(cell.value)}")
    if cells:
        print(f"行{row_idx}: {', '.join(cells)}")
